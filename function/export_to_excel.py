import os
import pandas as pd
import numpy as np
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ===== 파일 경로 =====
here = os.path.dirname(os.path.abspath(__file__))
result_dir = os.path.join(here, "..", "result")

# result 폴더에서 가장 최신 txt 선택
txt_files = [
    os.path.join(result_dir, f)
    for f in os.listdir(result_dir)
    if f.lower().endswith(".txt")
]
if not txt_files:
    raise FileNotFoundError(f"result 폴더에 txt 파일이 없습니다: {result_dir}")

grouped_txt_path = max(txt_files, key=os.path.getmtime)
pkl_path = os.path.join(here, "..", "pkl", "undefined.pkl")  # grouped.txt의 인덱스 범위에 맞는 데이터셋
excel_output_path = 'grouped_data.xlsx'

print(f"[INFO] 최신 grouped 파일 사용: {grouped_txt_path}")

# ===== grouped.txt 파일 읽기 =====
print("[1/4] grouped.txt 파일 읽는 중...")
with open(grouped_txt_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

# shape 정보 추출
shape_match = None
for line in lines:
    if 'shape=' in line:
        shape_match = re.search(r'shape=\((\d+),\s*(\d+)\)', line)
        break

if not shape_match:
    raise ValueError("grouped.txt 파일에서 shape 정보를 찾을 수 없습니다.")

rows = int(shape_match.group(1))
cols = int(shape_match.group(2))

# 각 행 파싱
groups = []
for line in lines:
    line = line.strip()
    if not line or line.startswith('created_at') or 'shape=' in line or 'dtype=' in line:
        continue
    # [숫자, 숫자, ...] 형태 파싱
    if line.startswith('['):
        # 마지막 ] 제거
        if line.endswith(','):
            line = line[:-1]
        if line.endswith(']'):
            line = line[:-1]
        if line.startswith('['):
            line = line[1:]
        # 쉼표로 분리하고 숫자 추출
        try:
            numbers = [int(x.strip()) for x in line.split(',') if x.strip()]
            if len(numbers) == cols:
                groups.append(numbers)
        except ValueError:
            continue

groups_array = np.array(groups, dtype=np.uint32)
print(f"  - 그룹 수: {len(groups_array)}개 (각 그룹당 {cols}개 인덱스)")

# ===== 원본 데이터 로드 =====
print("[2/4] 원본 데이터 로드 중...")
df_original = pd.read_pickle(pkl_path)
if isinstance(df_original, list):
    df_original = pd.DataFrame(df_original)
print(f"  - 원본 데이터 개수: {len(df_original)}개")

# ===== 엑셀 데이터 준비 =====
print("[3/4] 엑셀 데이터 준비 중...")
excel_data = []

for group_idx, group in enumerate(groups_array):
    for pos_in_group, data_idx in enumerate(group):
        if data_idx < len(df_original):
            row_data = {
                '그룹번호': group_idx + 1,
                '그룹내위치': pos_in_group + 1,
                '데이터인덱스': int(data_idx),
            }
            
            # 원본 데이터 정보 추가 (BATCH_20260129_161623_VPKO.pkl 규격)
            if 'chip_uid' in df_original.columns:
                row_data['chip_uid'] = df_original.iloc[data_idx]['chip_uid']
            if 'lot_name' in df_original.columns:
                row_data['lot_name'] = df_original.iloc[data_idx]['lot_name']
            if 'failure_type' in df_original.columns:
                row_data['failure_type'] = df_original.iloc[data_idx]['failure_type']
            if 'created_at' in df_original.columns:
                row_data['created_at'] = df_original.iloc[data_idx]['created_at']
            if 'tsv_status' in df_original.columns:
                row_data['tsv_status'] = df_original.iloc[data_idx]['tsv_status']
            if 'die_status' in df_original.columns:
                row_data['die_status'] = df_original.iloc[data_idx]['die_status']
            if 'tsv_coordinate' in df_original.columns:
                row_data['tsv_coordinate'] = df_original.iloc[data_idx]['tsv_coordinate']
            if 'coor_x' in df_original.columns:
                row_data['coor_x'] = df_original.iloc[data_idx]['coor_x']
            if 'coor_y' in df_original.columns:
                row_data['coor_y'] = df_original.iloc[data_idx]['coor_y']
            
            excel_data.append(row_data)
        else:
            # 범위를 벗어나는 인덱스는 건너뛰기
            pass

df_excel = pd.DataFrame(excel_data)
print(f"  - 엑셀 행 수: {len(df_excel)}개")

# ===== 엑셀 파일 저장 =====
print("[4/4] 엑셀 파일 저장 중...")
with pd.ExcelWriter(excel_output_path, engine='openpyxl') as writer:
    df_excel.to_excel(writer, sheet_name='Grouped Data', index=False)
    
    # 엑셀 스타일링
    workbook = writer.book
    worksheet = writer.sheets['Grouped Data']
    
    # 헤더 스타일
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    for col_num, column_title in enumerate(df_excel.columns, 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        # 컬럼 너비 자동 조정
        column_letter = get_column_letter(col_num)
        max_length = max(
            len(str(column_title)),
            df_excel[column_title].astype(str).str.len().max()
        )
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    # 데이터 정렬
    for row in range(2, len(df_excel) + 2):
        for col_num in range(1, len(df_excel.columns) + 1):
            cell = worksheet.cell(row=row, column=col_num)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 그룹별로 색상 구분 (선택사항)
    group_fills = [
        PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid"),
        PatternFill(start_color="FFF4E7", end_color="FFF4E7", fill_type="solid"),
    ]
    
    current_group = None
    for row in range(2, len(df_excel) + 2):
        group_num = df_excel.iloc[row-2]['그룹번호']
        if group_num != current_group:
            current_group = group_num
            fill = group_fills[group_num % 2]
        for col_num in range(1, len(df_excel.columns) + 1):
            worksheet.cell(row=row, column=col_num).fill = fill

print(f"[완료] 엑셀 파일 저장: {excel_output_path}")
print(f"   - 총 {len(groups_array)}개 그룹")
print(f"   - 총 {len(df_excel)}개 행")
if len(df_excel) < len(groups_array) * cols:
    print(f"   - 경고: 일부 인덱스가 데이터 범위를 벗어나 제외되었습니다.")

